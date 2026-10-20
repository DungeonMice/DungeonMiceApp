"""Demand-driven frame adapter. No changes or video outputs in the Python project."""
import collections
import hmac
import copy
import io
import json
import math
import os
from pathlib import Path
import re
import struct
import sys
import tempfile
import threading
import time
import zipfile
from experiment_config import select_config, validate_frame_regions

LOCK = threading.RLock()
SESSION = None
MAX_BYTES = 32 * 1024 * 1024
IDLE_SECONDS = 300


class Session:
    def __init__(self, identifier, request, options, normalize):
        self.id = identifier
        self.last = time.monotonic()
        self.cache = collections.OrderedDict()
        self.bytes = 0
        self.count = 0
        sys.path.insert(0, str(options.python_project / "src"))
        import cv2
        from input_loader import load_experiment
        from RunExperiment_per_frame import VideoAnalyzer
        source = Path(request["sourcePath"]).resolve(strict=True)
        self.source = source
        self.report_workspace = None
        self.report_packet = None
        self.completed = False
        if not source.is_file() or source.suffix.lower() not in {".mp4", ".mov", ".avi", ".mkv", ".m4v"}:
            raise ValueError("Video original no disponible")
        regions = normalize(request["regions"])
        if not regions:
            raise ValueError("Guarda las regiones antes de analizar")
        capture = cv2.VideoCapture(str(source))
        try:
            self.width = int(capture.get(cv2.CAP_PROP_FRAME_WIDTH))
            self.height = int(capture.get(cv2.CAP_PROP_FRAME_HEIGHT))
            if not capture.isOpened() or self.width <= 0 or self.height <= 0:
                raise ValueError("No se pudo abrir el video original")
            validate_frame_regions(regions, self.width, self.height)
            ok, self.first_frame = capture.read()
            if not ok:
                raise ValueError("No se pudo leer el primer frame")
        finally:
            capture.release()
        config = select_config(options, regions)
        self.kind = config["labyrinth_type"]
        config.update(video_path=str(source), regions=regions)
        with tempfile.TemporaryDirectory(prefix="dungeonmice-frame-config-") as temp:
            path = Path(temp) / "experiment.json"
            path.write_text(json.dumps(config), encoding="utf-8")
            labyrinth = load_experiment(path)
            self.analyzer = VideoAnalyzer(labyrinth)
        if self.analyzer.fps <= 0 or self.analyzer.total_frames <= 0:
            self.analyzer.release()
            raise ValueError("FPS o duracion invalidos")
        self.cv2 = cv2
        self.views = {}
        original = self.analyzer._visualizer.draw_trajectory_on_mask
        def capture_views(mask, *args, **kwargs):
            self.views["detection"] = mask.copy()
            rendered = original(mask, *args, **kwargs)
            self.views["trajectory"] = rendered
            return rendered
        self.analyzer._visualizer.draw_trajectory_on_mask = capture_views

    def info(self):
        return dict(sessionId=self.id, width=self.width, height=self.height,
                    fps=self.analyzer.fps, total=self.analyzer.total_frames,
                    processed=self.count, firstCached=next(iter(self.cache), 0), cachedBytes=self.bytes,
                    experiment=self.kind, views=["video", "detection", "trajectory"],
                    engine=type(self.analyzer).__module__ + "." + type(self.analyzer).__name__)

    def next(self):
        data = self.analyzer.next()
        if data is None:
            self.completed = self.count == self.analyzer.total_frames
            return None
        # The pulled engine otherwise retains every full-resolution BGR image.
        # Keep tracker state intact, but own only a bounded compressed preview window.
        self.analyzer._cache.clear()
        images = dict(video=data.image, **getattr(self, "views", {}))
        payload = bytearray()
        descriptors = []
        for name, image in images.items():
            ok, encoded = self.cv2.imencode(".jpg", image, [self.cv2.IMWRITE_JPEG_QUALITY, 88])
            if not ok:
                raise ValueError("No se pudo codificar la vista " + name)
            raw = encoded.tobytes()
            descriptors.append(dict(id=name, offset=len(payload), length=len(raw)))
            payload.extend(raw)
        states = {name: dict(inside=bool(state.inside), entries=int(state.entries),
                            totalTime=float(state.total_time)) for name, state in data.region_states.items()}
        position = None if data.position is None else [int(value) for value in data.position]
        header = json.dumps(dict(index=data.index, time=float(data.time), position=position,
                                regions=states, recordingActive=bool(data.recording_active), images=descriptors)).encode("utf-8")
        packet = struct.pack(">I", len(header)) + header + payload
        if len(packet) > MAX_BYTES:
            raise ValueError("Frame demasiado grande")
        self.count = data.index
        self.cache[data.index] = packet
        self.bytes += len(packet)
        while len(self.cache) > 120 or self.bytes > MAX_BYTES:
            _, old = self.cache.popitem(last=False)
            self.bytes -= len(old)
        return packet

    def reports(self):
        if not self.completed:
            raise ValueError("Los reportes estan disponibles al finalizar el analisis completo")
        if self.report_packet is not None:
            return self.report_packet
        from openpyxl import load_workbook
        self.report_workspace = tempfile.TemporaryDirectory(prefix="dungeonmice-reports-")
        folder = Path(self.report_workspace.name)
        labyrinth = self.analyzer._labyrinth
        name = re.sub(r'[\[\]:*?/\\]', '_', self.source.stem)[:31] or "video"
        # The original writer derives output paths from this map. Use an isolated
        # directory, never the source video directory or a process-wide chdir.
        labyrinth.process_video(
            all_results={name: copy.deepcopy(self.analyzer._logic.states)},
            all_trajectories={name: (labyrinth.trajectory_x, labyrinth.trajectory_y)},
            all_video_paths={name: str(folder / self.source.name)},
            all_first_frames={name: self.first_frame},
            all_start_times={name: labyrinth.start_time},
            all_recording_durations={name: max(0, self.count / self.analyzer.fps - labyrinth.start_time)})
        files = sorted(folder.glob("*.xlsx"))
        if len(files) != 1:
            raise ValueError("Python no genero el reporte Excel esperado")
        workbook = load_workbook(files[0], data_only=True, read_only=True)
        try:
            sheets = []
            for sheet in workbook:
                rows = [[str(cell) if cell is not None else "" for cell in row] for row in sheet.iter_rows(values_only=True)]
                sheets.append(dict(name=sheet.title, rows=rows))
        finally:
            workbook.close()
        metadata = dict(version=1, sourceName=self.source.name, frames=self.count,
                        fps=self.analyzer.fps, experiment=self.kind, workbook=files[0].name, sheets=sheets)
        (folder / "report.json").write_text(json.dumps(metadata, ensure_ascii=False), encoding="utf-8")
        packet = io.BytesIO()
        with zipfile.ZipFile(packet, "w", zipfile.ZIP_DEFLATED) as archive:
            for file in sorted(folder.iterdir()):
                if file.suffix in {".xlsx", ".png", ".json"}:
                    archive.write(file, file.name)
        self.report_packet = packet.getvalue()
        return self.report_packet

    def close(self):
        try:
            self.analyzer.release()
        finally:
            self.cache.clear()
            self.bytes = 0
            if getattr(self, "report_workspace", None) is not None:
                self.report_workspace.cleanup()


def close_all(slot):
    global SESSION
    with LOCK:
        if SESSION is not None:
            try:
                SESSION.close()
            finally:
                SESSION = None
                slot.release()


def expire(slot, stopping):
    while not stopping.wait(15):
        with LOCK:
            if SESSION and time.monotonic() - SESSION.last > IDLE_SECONDS:
                close_all(slot)


def handle(handler, options, normalize, slot, stopping):
    global SESSION
    if not handler.path.startswith("/api/v1/frame-sessions"):
        return False
    token = os.environ.get("DUNGEONMICE_SHUTDOWN_TOKEN", "")
    if not token or not hmac.compare_digest(token, handler.headers.get("X-DungeonMice-Token", "")):
        handler.reply(403, {"error": "Frame API requires the local application token"})
        return True
    match = re.fullmatch(r"/api/v1/frame-sessions/([a-f0-9]{32})(?:/(next|reports|[0-9]+))?", handler.path)
    try:
        handler.connection.settimeout(30)
        with LOCK:
            if stopping.is_set():
                handler.reply(503, {"error": "Service stopping"})
            elif handler.command == "POST" and handler.path == "/api/v1/frame-sessions":
                handler.connection.settimeout(30)
                raw = handler.rfile.read(handler.length(65536))
                request = json.loads(raw)
                identifier = request.get("sessionId", "")
                if not isinstance(identifier, str) or not re.fullmatch(r"[a-f0-9]{32}", identifier):
                    raise ValueError("Invalid session ID")
                if SESSION is not None or not slot.acquire(blocking=False):
                    handler.reply(409, {"error": "Ya hay un analisis activo"})
                else:
                    try:
                        SESSION = Session(identifier, request, options, normalize)
                    except Exception:
                        slot.release()
                        raise
                    handler.reply(201, SESSION.info())
            elif not match or SESSION is None or match.group(1) != SESSION.id:
                handler.reply(410, {"error": "Sesion no disponible"})
            elif handler.command == "DELETE" and match.group(2) is None:
                close_all(slot)
                handler.reply(200, {"closed": True})
            elif handler.command == "GET":
                SESSION.last = time.monotonic()
                action = match.group(2)
                if action is None:
                    handler.reply(200, SESSION.info())
                elif action == "reports":
                    packet = SESSION.reports()
                    handler.send_response(200)
                    handler.send_header("Content-Type", "application/zip")
                    handler.send_header("Content-Length", str(len(packet)))
                    handler.send_header("Connection", "close")
                    handler.end_headers()
                    handler.wfile.write(packet)
                    handler.close_connection = True
                else:
                    packet = SESSION.next() if action == "next" else SESSION.cache.get(int(action))
                    if packet is None:
                        if action == "next":
                            handler.send_response(204)
                            handler.send_header("Connection", "close")
                            handler.end_headers()
                            handler.close_connection = True
                        else:
                            handler.reply(410, {"error": "Frame fuera de la cache disponible"})
                    else:
                        length = struct.unpack(">I", packet[:4])[0]
                        metadata = json.loads(packet[4:4 + length])
                        metadata["firstCached"] = next(iter(SESSION.cache))
                        header = json.dumps(metadata).encode("utf-8")
                        packet = struct.pack(">I", len(header)) + header + packet[4 + length:]
                        handler.send_response(200)
                        handler.send_header("Content-Type", "application/vnd.dungeonmice.frame")
                        handler.send_header("Content-Length", str(len(packet)))
                        handler.send_header("Connection", "close")
                        handler.end_headers()
                        handler.wfile.write(packet)
                        handler.close_connection = True
            else:
                handler.reply(405, {"error": "Method not allowed"})
    except (BrokenPipeError, ConnectionResetError):
        close_all(slot)
    except Exception as error:
        handler.reply(422, {"error": str(error)[:600]})
    return True
